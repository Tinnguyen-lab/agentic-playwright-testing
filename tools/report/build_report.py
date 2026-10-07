"""Dựng Báo cáo chuyên đề ĐATN (.docx) trên nền file mẫu của nhóm.

    python build_report.py <rq2_results_both.json> <rq2_results_both_v1.json> <out.docx>

Số liệu RQ2 đọc thẳng từ JSON kết quả; RQ1 lấy từ docs/development/eval-results.md (đã commit).
"""
from __future__ import annotations

import json
import math
import os
import sys
from collections import Counter
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH

from report_lib import Report

HERE = Path(__file__).parent
ROOT = HERE.parents[1]
TEMPLATE = os.getenv("REPORT_TEMPLATE", r"E:\OneDrive\IT Learning\HK4\Cyber security\DoAn_CyberSecurity\doc\Báo cáo đồ án.docx")
SITE_NAME = {
    "saucedemo": "SauceDemo", "the-internet": "The Internet (Heroku)", "uitestingplayground": "UI Testing Playground",
    "practicetestautomation": "Practice Test Automation", "demoqa": "DemoQA", "orangehrm": "OrangeHRM Demo",
}


def pct(a, b):
    return f"{a / b * 100:.0f}%" if b else "—"


def mcnemar_exact_p(b: int, c: int) -> float:
    """McNemar chính xác (nhị thức hai phía) trên các cặp bất đồng b, c."""
    n = b + c
    if n == 0:
        return 1.0
    k = min(b, c)
    tail = sum(math.comb(n, i) for i in range(k + 1)) / 2 ** n
    return min(1.0, 2 * tail)


def front_matter(r: Report):
    r.set_cover_line("BÁO CÁO ĐỒ ÁN", "BÁO CÁO CHUYÊN ĐỀ")
    r.set_cover_line("MÔN NHẬP MÔN", "ĐỒ ÁN TỐT NGHIỆP")
    r.set_cover_line("Đề tài:", "Đề tài: HỆ THỐNG AGENTIC AI BÁN TỰ ĐỘNG SINH VÀ KIỂM CHỨNG KIỂM THỬ ĐẦU-CUỐI DỰA TRÊN PLAYWRIGHT")
    r.set_cover_line("Kết hợp Nikto", "Từ tài liệu yêu cầu đến mã Playwright chạy được,")
    r.set_cover_line("để phân tích", "có truy vết, có neo DOM và có con người phê duyệt")
    r.set_cover_line("GVHD:", "GVHD: [Họ tên giảng viên hướng dẫn]")
    r.set_cover_line("1. Võ Quỳnh Nhi", "1. Nguyễn Lê Sang").runs[0].text = "1. Nguyễn Lê Sang"
    m1 = next(p for p in r.doc.paragraphs if p.text.startswith("1. Nguyễn Lê Sang"))
    m1.add_run("\tMSSV: 25210178")
    m2 = r.set_cover_line("2. Nguyễn Lê Sang", "2. [Họ tên thành viên 2]")
    m2.add_run("\tMSSV: [...]")
    for start in ("3. Dương Quốc Cường", "4. Nguyễn Trung Tín"):
        p = next(p for p in r.doc.paragraphs if p.text.startswith(start))
        p._p.getparent().remove(p._p)
    r.set_cover_line("Tp. Hồ Chí Minh", "Tp. Hồ Chí Minh, 10/2026")
    r.set_header("Báo cáo chuyên đề đồ án tốt nghiệp")

    r.h1("NHẬN XÉT CỦA GIÁO VIÊN HƯỚNG DẪN")
    r.dotted_lines(18)

    r.h1("LỜI MỞ ĐẦU")
    r.p("Kiểm thử đầu-cuối (end-to-end, E2E) kiểm tra một ứng dụng web theo đúng cách người dùng thật sử dụng nó: mở trang, "
        "nhập liệu, bấm nút và quan sát kết quả. Viết và duy trì các kịch bản này tốn nhiều công. Kiểm thử viên phải đọc tài "
        "liệu yêu cầu, tự suy ra test case, rồi viết mã tự động hoá và sửa lại mỗi khi giao diện thay đổi.")
    r.p("Mô hình ngôn ngữ lớn (LLM) hiện đọc được tài liệu yêu cầu và viết được mã kiểm thử. Khi đưa vào quy trình thật, "
        "chúng gặp bốn vấn đề: tài liệu yêu cầu thường mơ hồ nhưng LLM vẫn tự điền chỗ trống; không còn dấu vết nối test với "
        "yêu cầu gốc; locator do LLM đoán có thể không tồn tại trên trang; và cơ chế tự sửa lỗi có thể làm yếu assertion để "
        "test \"xanh\" giả. Đồ án này xây dựng một hệ thống Agentic AI bán tự động cho chuỗi từ tài liệu yêu cầu đến mã "
        "Playwright Python chạy được. Trong hệ thống, các agent chỉ phân tích, sinh và đề xuất. Con người giữ quyền quyết "
        "định ở các cổng ảnh hưởng tới ý nghĩa kiểm thử.")
    r.p("Nội dung báo cáo gồm sáu chương:")
    r.bullets(["Chương 1: Tổng quan đề tài", "Chương 2: Cơ sở lý thuyết và nghiên cứu liên quan",
               "Chương 3: Phân tích, thiết kế và xây dựng hệ thống",
               "Chương 4: Môi trường thực nghiệm và phương pháp đánh giá", "Chương 5: Kết quả thực nghiệm và đánh giá",
               "Chương 6: Kết luận và hướng phát triển"])
    r.p("Toàn bộ thực nghiệm chỉ chạy trên các website demo công khai được dựng ra để luyện kiểm thử tự động (SauceDemo, "
        "The Internet, UI Testing Playground, DemoQA, Practice Test Automation, OrangeHRM Demo). Mọi con số trong báo cáo "
        "đều tái lập được bằng lệnh ghi trong Chương 4.")

    r.h1("BẢNG PHÂN CÔNG CÔNG VIỆC")
    r.table("Bảng phân công công việc", ["Thành viên", "Vai trò", "Công việc chính"], [
        ["Nguyễn Lê Sang", "AI / Web / Data", "Kiến trúc agent, tích hợp LLM (cloud và local), Requirement Agent, bộ đánh giá RQ1, "
         "giao diện Streamlit, thực nghiệm RQ2"],
        ["[Họ tên thành viên 2]", "QA / Playwright", "Khảo sát Playwright và locator, tiêu chí đánh giá test case, yêu cầu cho "
         "traceability và constrained repair, catalog target RQ2"],
        ["Cả nhóm", "", "Chốt research gap và MVP, review tài liệu và pull request, ghi biên bản họp"],
    ], [3.5, 3, 9.5], prefix="0.1")

    r.h1("DANH MỤC THUẬT NGỮ")
    r.table("Bảng danh mục thuật ngữ", ["Thuật ngữ", "Giải thích"], [
        ["Kiểm thử đầu-cuối (E2E)", "Kiểm thử toàn bộ luồng nghiệp vụ qua giao diện như người dùng thật"],
        ["Agent (tác tử)", "Thành phần dùng LLM để thực hiện một nhiệm vụ có đầu vào, đầu ra theo schema xác định"],
        ["Locator", "Biểu thức xác định phần tử trên trang, ví dụ get_by_role(\"button\", name=\"Login\")"],
        ["Grounding (neo DOM)", "Kiểm chứng locator trên trang thật: phần tử có tồn tại và khớp đúng một hay không"],
        ["Accessibility tree", "Cây trợ năng trình duyệt dựng từ DOM: mỗi nút có role, tên hiển thị và trạng thái"],
        ["ARIA snapshot", "Bản chụp cây trợ năng dạng YAML do Playwright sinh (locator.aria_snapshot())"],
        ["Test oracle", "Cơ chế quyết định test đạt hay không; ở đây là các assertion expect_*"],
        ["Self-healing", "Tự sửa locator khi phần tử đổi thuộc tính mà nghiệp vụ không đổi"],
        ["Human-in-the-loop", "Con người duyệt tại các điểm quyết định thay vì để hệ thống tự vận hành"],
        ["Boundary Value Analysis", "Kỹ thuật thiết kế test chọn giá trị tại và ngay ngoài biên của ràng buộc"],
        ["First-run pass rate", "Tỉ lệ script do LLM sinh chạy đạt ngay lần đầu, không qua sửa"],
        ["Over-flag", "Tỉ lệ yêu cầu sạch bị gắn cờ mơ hồ nhầm"],
        ["Assertion weakening", "Làm yếu lệnh kiểm chứng (ví dụ toBe(5) thành toBeTruthy()) để test đạt giả"],
    ], [4.5, 11.5], prefix="0.2")

    r.h1("DANH MỤC TỪ VIẾT TẮT")
    r.table("Bảng danh mục từ viết tắt", ["Từ viết tắt", "Viết đầy đủ", "Ý nghĩa"], [
        ["AG", "Approval Gate", "Cổng phê duyệt của con người"],
        ["ARIA", "Accessible Rich Internet Applications", "Chuẩn W3C mô tả vai trò, tên của phần tử cho công nghệ hỗ trợ"],
        ["BVA", "Boundary Value Analysis", "Phân tích giá trị biên"],
        ["DOM", "Document Object Model", "Cây đối tượng của trang HTML"],
        ["DSL", "Domain-Specific Language", "Tập action Playwright rút gọn mà LLM được phép sinh"],
        ["E2E", "End-to-End", "Kiểm thử đầu-cuối"],
        ["HITL", "Human-in-the-loop", "Con người trong vòng lặp"],
        ["LLM", "Large Language Model", "Mô hình ngôn ngữ lớn"],
        ["RQ", "Research Question", "Câu hỏi nghiên cứu"],
        ["SRS", "Software Requirements Specification", "Đặc tả yêu cầu phần mềm"],
        ["TC", "Test Case", "Ca kiểm thử"],
        ["UI", "User Interface", "Giao diện người dùng"],
    ], [2.6, 5.4, 8], prefix="0.3")

    r.h1("DANH MỤC CÁC BẢNG, HÌNH ẢNH")
    r.p("Danh mục các bảng:", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    r.field('TOC \\h \\z \\t "CapBang,1"', "Mở bằng Word và chọn Update Field (F9) để tạo danh mục bảng.")
    r.p("Danh mục các hình:", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    r.field('TOC \\h \\z \\t "CapHinh,1"', "Mở bằng Word và chọn Update Field (F9) để tạo danh mục hình.")

    r.page_break()
    r.doc.add_paragraph("MỤC LỤC", style="TocTitle")
    r.field('TOC \\o "1-3" \\h \\z \\u', "Mở bằng Word và chọn Update Field (F9) để tạo mục lục.")


def chapter1(r: Report):
    r.h1("Chương 1: TỔNG QUAN ĐỀ TÀI", chapter=True)
    r.h2("1.1 Đặt vấn đề và lý do chọn đề tài")
    r.p("Trong quy trình phát triển phần mềm, kiểm thử E2E là chốt chặn cuối trước khi phát hành. Việc chuyển thủ công tài "
        "liệu yêu cầu (SRS, use case, user story) thành kịch bản và mã kiểm thử tốn thời gian và khó bảo trì. Một nghiên cứu "
        "công nghiệp quy tới 73% thất bại của test suite về lỗi locator chứ không phải lỗi chức năng [6]. Các nghiên cứu "
        "2024–2026 cho thấy LLM sinh được test case từ tài liệu yêu cầu [5][7][11] và thao tác được trên web thật [10]. Khi "
        "áp dụng thực tế, nhóm nhận thấy bốn vấn đề mà chưa nguồn nào trong 12 nguồn khảo sát giải quyết trọn vẹn:")
    r.bullets([
        "Yêu cầu chưa được quản trị như một artifact có cấu trúc. Các công cụ thường đưa nguyên tài liệu cho LLM, giả định "
        "đầu vào đầy đủ và không có bước phát hiện yêu cầu thiếu, mơ hồ hoặc mâu thuẫn [4][7]. LLM vì thế tự điền giả định "
        "vào expected result.",
        "Truy vết bị đứt. Ma trận truy vết thường nằm trong Excel hoặc hệ thống ALM tách khỏi mã, và khoảng 80% người hành nghề "
        "coi chi phí cập nhật thủ công là rào cản chính [9]. Khi LLM tự sinh test, không ai biết đoạn mã nào phục vụ yêu cầu nào.",
        "Locator do LLM đoán có thể không tồn tại. Không neo vào DOM thật, mô hình suy ra selector từ mẫu đã thấy khi huấn luyện; "
        "một case study ghi nhận agent bỏ qua công cụ kiểm chứng DOM và gây hàng loạt lỗi timeout [8].",
        "Tự sửa lỗi không ràng buộc làm hỏng oracle. Trên 300 báo cáo thực thi, một hệ multi-agent chỉ hội tụ thật khoảng 50% "
        "thay vì 70% bề mặt, vì agent làm yếu assertion hoặc xoá test đang lỗi để đạt trạng thái Pass [8].",
    ])
    r.p("Đề tài chọn hướng bán tự động: LLM làm phần nặng (phân tích, sinh, đề xuất), còn các quyết định làm thay đổi ý nghĩa "
        "kiểm thử phải qua con người. Hướng này khớp với khuyến nghị của [8] về giới hạn vòng sửa, phát hiện trôi ngữ nghĩa và "
        "bắt buộc phê duyệt thay đổi assertion.")

    r.h2("1.2 Mục tiêu của đề tài")
    r.bullets([
        "Xây dựng pipeline năm agent: Requirement Analysis, Test Design, Playwright Generation, Execution và Repair, nối từ tài "
        "liệu yêu cầu tới script Playwright Python chạy trên website thật.",
        "Chuẩn hoá yêu cầu thành schema có cấu trúc và gắn cờ mơ hồ theo rubric, chặn ở cổng duyệt trước khi sinh test.",
        "Duy trì truy vết từ yêu cầu tới test case, script, kết quả chạy và đề xuất sửa.",
        "Giới hạn tự sửa lỗi bằng chính sách rủi ro tất định (không do LLM tự đánh giá) và ngân sách sửa.",
        "Đo định lượng: độ chính xác phát hiện mơ hồ (RQ1) và tỉ lệ script chạy đạt ngay lần đầu khi có và không có ngữ cảnh "
        "DOM (RQ2).",
    ])

    r.h2("1.3 Câu hỏi nghiên cứu")
    r.table("Bốn câu hỏi nghiên cứu và trạng thái trong chuyên đề", ["RQ", "Câu hỏi", "Trạng thái"], [
        ["RQ1", "Chuẩn hoá tài liệu yêu cầu thành biểu diễn có cấu trúc có cải thiện tính hợp lệ và độ bao phủ của test case "
         "so với đưa thẳng tài liệu cho LLM không?", "Đã đo: 15 tài liệu, 47 yêu cầu, 83 điều kiện gold"],
        ["RQ2", "Kết hợp yêu cầu với DOM hoặc accessibility tree có cải thiện tỉ lệ mã Playwright chạy đạt ở lần đầu không?",
         "Đã đo: 50 target, 6 site, 3 nhánh"],
        ["RQ3", "Sửa lỗi có ràng buộc và phê duyệt của con người có tăng tỉ lệ sửa thành công mà không đổi ý nghĩa test case "
         "không?", "Đã đo: ShopLab, 14 mutation, 3 nhánh"],
        ["RQ4", "Hệ thống giảm bao nhiêu thời gian và thao tác thủ công so với quy trình truyền thống?",
         RQ4_STATUS],
    ], [1.3, 10.7, 4], center_cols=(0,))

    r.h2("1.4 Phạm vi nghiên cứu")
    r.h3("1.4.1 Trong phạm vi")
    r.p("Đề tài nhắm vào kiểm thử chức năng E2E cho ứng dụng web, sinh mã Playwright bằng Python. Đầu vào là tài liệu DOCX, "
        "PDF dạng văn bản, Markdown hoặc text, chứa use case, user story kèm acceptance criteria hoặc functional requirement. "
        "Hệ thống sinh test case positive, negative, boundary, error guessing và alternative flow; chạy test, thu bằng chứng; "
        "đề xuất sửa trong ràng buộc và có con người phê duyệt.")
    r.h3("1.4.2 Giới hạn của đề tài")
    r.bullets([
        "Không kiểm thử mobile, desktop, hiệu năng, bảo mật hay accessibility toàn diện.",
        "Không sinh đồng thời Selenium, Cypress và Playwright; chỉ Playwright Python, trình duyệt Chromium, chạy tuần tự.",
        "Không tự sửa mã nguồn của ứng dụng được kiểm thử và không tự suy ra oracle từ yêu cầu mơ hồ.",
        "Không huấn luyện hay fine-tune LLM; chỉ dùng mô hình có sẵn qua API tương thích OpenAI.",
        "DSL chỉ gồm goto, fill, click và ba loại assertion; kịch bản cần chọn dropdown, hover, nhấn phím hay double click "
        "nằm ngoài phạm vi.",
        "Thực nghiệm dùng ứng dụng demo và ứng dụng ShopLab do nhóm tự dựng, không phải ứng dụng doanh nghiệp.",
    ])

    r.h2("1.5 Đóng góp chính")
    r.p("Đề tài không tuyên bố phát minh từng thành phần riêng lẻ. Phần mới nằm ở cách ghép và kiểm soát chúng trong một "
        "pipeline Playwright Python:")
    r.bullets([
        "Requirement Analysis Agent với rubric bảy loại mơ hồ, chỉ phát hiện mà không tự giải quyết, chờ duyệt ở AG-01. Cột "
        "\"ambiguity detection\" trống trên cả 12 nguồn khảo sát.",
        "Sinh test biên và âm bản tất định bằng Boundary Value Analysis từ ràng buộc số, bổ sung cho kết quả của LLM.",
        "Grounding đa bước: đi theo luồng và kiểm chứng locator trên DOM tại đúng bước đó, bắt được phần tử chỉ xuất hiện sau "
        "điều hướng.",
        "Repair policy tất định phân mức rủi ro theo nội dung thay đổi (locator, dữ liệu, assertion, bước), ngân sách sửa bằng "
        "2 và self-healing locator từ DOM sống, mọi đề xuất đều cần người duyệt.",
        "Bộ đánh giá có ground truth: RQ1 trên 47 yêu cầu chèn khuyết tật có chủ đích và 83 điều kiện kiểm thử gold; RQ2 "
        "trên catalog 50 target có oracle viết tay chứng minh mỗi target làm được; RQ3 trên ShopLab, một ứng dụng đích có "
        "14 mutation gắn nhãn kỹ thuật/ngữ nghĩa, đo được việc cơ chế sửa có che lỗi thật hay không.",
        "Hệ thống chạy được trọn luồng trên giao diện, lưu version, quyết định duyệt append-only và chuỗi truy vết vào DB.",
    ])

    r.h2("1.6 Công nghệ sử dụng")
    r.table("Công nghệ sử dụng", ["Thành phần", "Công nghệ", "Vai trò"], [
        ["Ngôn ngữ", "Python 3.14", "Toàn bộ hệ thống và script kiểm thử sinh ra"],
        ["Tự động hoá trình duyệt", "Playwright 1.62 (sync API), Chromium", "Grounding, aria snapshot, chạy script"],
        ["Mô hình dữ liệu", "Pydantic 2", "Schema đầu ra của agent, validate JSON"],
        ["LLM cloud", "DeepSeek (deepseek-chat) qua API tương thích OpenAI", "Backend chính cho thực nghiệm"],
        ["LLM local", "LM Studio + google/gemma-4-12b (GPU 8 GB)", "Backend miễn phí, chạy offline"],
        ["Sinh mã", "Jinja2", "Render PlaywrightPlan thành script Python tự chứa"],
        ["Đọc tài liệu", "python-docx, PyMuPDF", "Trích văn bản DOCX/PDF, giữ mốc trang để truy vết"],
        ["Giao diện", "Streamlit", "Quy trình duyệt AG-01..05, sinh và chạy script, lịch sử"],
        ["Lưu trữ", "SQLAlchemy 2; SQLite / SQL Server Express", "Version, quyết định, truy vết, nhật ký thao tác"],
        ["Ứng dụng đích", "Flask (ShopLab)", "App có mutation kiểm soát được cho RQ3, RQ4"],
        ["Kiểm thử", "pytest, pytest-playwright", "Unit test offline; chạy test của người tham gia RQ4"],
    ], [3.6, 6, 6.4])


def chapter2(r: Report):
    r.h1("Chương 2: CƠ SỞ LÝ THUYẾT VÀ NGHIÊN CỨU LIÊN QUAN", chapter=True)
    r.h2("2.1 Kiểm thử đầu-cuối và Playwright")
    r.h3("2.1.1 Kiểm thử đầu-cuối")
    r.p("Kiểm thử E2E chạy ứng dụng hoàn chỉnh qua giao diện, mô phỏng hành vi người dùng từ đầu tới cuối một luồng nghiệp vụ. "
        "Loại kiểm thử này bắt được lỗi tích hợp mà unit test bỏ sót, đổi lại chậm và dễ gãy khi giao diện thay đổi. Mỗi test "
        "gồm chuỗi bước (mở trang, nhập, bấm) và một hoặc nhiều assertion đóng vai oracle.")
    r.h3("2.1.2 Playwright và chính sách locator")
    r.p("Playwright là thư viện tự động hoá trình duyệt của Microsoft, hỗ trợ Chromium, Firefox và WebKit. Hai cơ chế giúp "
        "test ổn định hơn. Thứ nhất là auto-wait: thao tác chờ phần tử sẵn sàng, assertion expect() tự thử lại đến hết timeout. "
        "Thứ hai là locator hướng người dùng [2]. Tài liệu chính thức khuyến nghị ưu tiên get_by_role, get_by_label, "
        "get_by_placeholder, get_by_text và get_by_test_id, hạn chế CSS/XPath vì chúng bám vào cấu trúc DOM dễ đổi. Locator "
        "ở chế độ strict: nếu khớp nhiều hơn một phần tử, thao tác báo lỗi. Hệ thống tận dụng điều này để đo chất lượng "
        "locator bằng số phần tử khớp.")
    r.p("Từ 2025, Playwright có bộ Test Agents gồm Planner, Generator và Healer [1]: Planner khám phá ứng dụng và viết kế "
        "hoạch Markdown, Generator chuyển thành test, Healer chạy lại và vá test hỏng. Bộ này là baseline sản phẩm gần nhất của "
        "đề tài nhưng không có bước phát hiện mơ hồ, không có cổng duyệt bắt buộc và chưa công bố chính sách sửa chi tiết.")
    r.h3("2.1.3 Accessibility tree và ARIA snapshot")
    r.p("Trình duyệt dựng accessibility tree từ DOM cho công nghệ hỗ trợ (trình đọc màn hình). Mỗi nút có role (button, "
        "textbox, link, heading...), tên hiển thị (accessible name) và trạng thái. Cây này gọn hơn DOM nhiều lần vì bỏ các "
        "thẻ trình bày, và nó dùng đúng ngôn ngữ của get_by_role. Playwright từ bản 1.49 có locator.aria_snapshot() trả cây "
        "này dạng YAML, ví dụ:")
    r.code('- heading "Dynamic Controls" [level=4]\n- checkbox\n- text: A checkbox\n- button "Remove"\n- textbox [disabled]\n- button "Enable"')
    r.p("Đây là ngữ cảnh DOM mà nhánh DOM-aware của RQ2 đưa cho LLM. Hướng a11y-first cũng là lõi của [6], còn [10] lọc DOM "
        "rồi đánh chỉ số phần tử cho LLM chọn.")

    r.h2("2.2 Mô hình ngôn ngữ lớn và hệ tác tử")
    r.h3("2.2.1 Structured output")
    r.p("LLM sinh văn bản tự do, trong khi pipeline cần dữ liệu có cấu trúc để validate và lưu. Các API tương thích OpenAI hỗ "
        "trợ response_format kiểu json_schema: máy chủ ép đầu ra theo ngữ pháp sinh từ schema, nên mô hình không thể trả sai "
        "kiểu hay thiếu trường. Hệ thống khai báo schema bằng Pydantic, gửi kèm response_format và validate lại phía client.")
    r.h3("2.2.2 Agent và con người trong vòng lặp")
    r.p("Trong đồ án, agent là một bước có đầu vào và đầu ra theo schema, dùng LLM ở phần cần suy luận và dùng code tất định "
        "ở phần còn lại. Mô hình human-in-the-loop đặt con người tại các điểm có rủi ro ngữ nghĩa: duyệt yêu cầu đã chuẩn hoá, "
        "duyệt test case, duyệt mọi đề xuất sửa đổi assertion, expected result hay bước kiểm thử. Bằng chứng ở [8] cho thấy "
        "thiếu các cổng này thì agent tối ưu cho trạng thái Pass thay vì cho tính đúng.")

    r.h2("2.3 Kỹ thuật thiết kế test case")
    r.p("Kiểm thử hộp đen dựa trên đặc tả có các kỹ thuật kinh điển trong chuẩn ISO/IEC/IEEE 29119-4: phân hoạch tương đương, "
        "phân tích giá trị biên, đoán lỗi và kiểm thử luồng thay thế. Với ràng buộc \"mật khẩu 8–20 ký tự\", BVA chọn các giá "
        "trị 8 và 20 (biên hợp lệ) cùng 7 và 21 (ngay ngoài biên, ca âm bản). [7] dùng các kỹ thuật này trong prompt; đồ án "
        "tách BVA ra thành code tất định để bảo đảm luôn có ca biên, không phụ thuộc LLM nhớ ra.")

    r.h2("2.4 Truy vết yêu cầu")
    r.p("Requirements traceability là khả năng đi xuôi từ một yêu cầu tới các artifact hiện thực nó và đi ngược lại. ReqToCode "
        "[9] nhúng phần tử truy vết vào mã để lỗi truy vết thành lỗi biên dịch, nhưng chỉ bảo đảm sự hiện diện cấu trúc chứ "
        "không bảo đảm đúng ngữ nghĩa, và không bao phủ kết quả chạy hay đề xuất sửa. Đồ án dùng chuỗi ID tất định "
        "REQ-001 → REQ-001-TC-01 → script → kết quả → đề xuất sửa để tính độ phủ và tìm artifact mồ côi.")

    r.h2("2.5 Self-healing và bài toán oracle")
    r.p("Self-healing locator tìm phần tử thay thế khi locator cũ gãy. Similo và HybridSimilo [12] chấm điểm ứng viên trên "
        "nhiều thuộc tính (tag, id, class, text, aria-label, vị trí) và đạt tỉ lệ định vị lại cao trên benchmark của họ. [6] "
        "chỉ trích lại đúng selector hỏng từ accessibility tree, không tốn lời gọi API. Các kỹ thuật này sửa được phần kỹ thuật "
        "nhưng không trả lời được câu hỏi khi nào thay đổi là hợp lệ về nghiệp vụ. Khi agent được phép sửa cả assertion, nó có "
        "thể sửa oracle cho khớp với lỗi [8]. Ranh giới mà đồ án chọn: locator tương đương là thay đổi kỹ thuật; assertion, "
        "expected result và bước kiểm thử là thay đổi ngữ nghĩa và phải qua người duyệt.")

    r.h2("2.6 Chỉ số đánh giá")
    r.bullets([
        "Precision = TP / (TP + FP): trong các cờ mơ hồ hệ thống gắn, bao nhiêu phần đúng. Recall = TP / (TP + FN): trong các "
        "khuyết tật thật, hệ thống bắt được bao nhiêu. F1 là trung bình điều hoà của hai chỉ số.",
        "Micro-F1 gộp TP/FP/FN của mọi loại rồi tính; macro-F1 tính F1 từng loại rồi lấy trung bình, nên loại hiếm có trọng "
        "số ngang loại phổ biến.",
        "Over-flag: tỉ lệ yêu cầu sạch (không có khuyết tật chèn) bị gắn cờ. Chỉ số này đo gánh nặng duyệt nhầm cho con người.",
        "First-run pass rate: số script đạt ngay lần chạy đầu chia cho số target. Script không được sửa tay hay qua Repair Agent.",
        "Kiểm định McNemar: hai nhánh chạy trên cùng target nên kết quả là dữ liệu cặp. McNemar chỉ dùng các cặp bất đồng "
        "(nhánh này đạt, nhánh kia trượt); với cỡ mẫu nhỏ dùng dạng chính xác theo phân phối nhị thức.",
    ])

    r.h2("2.7 Các nghiên cứu liên quan")
    r.p("Nhóm khảo sát 12 nguồn chia năm mảng: baseline sản phẩm, requirement → test bằng LLM, traceability, DOM grounding và "
        "self-healing có ràng buộc. Bảng 2.1 tóm tắt năng lực của các nguồn theo tám tiêu chí mà đề tài cần.")
    Y, H, N = "Có", "Một phần", "Không"
    r.table("Đối chiếu năng lực 12 nguồn với nhu cầu đề tài", ["Nguồn", "Yêu cầu cấu trúc", "Phát hiện mơ hồ", "Truy vết",
                                                             "Neo DOM", "Oracle", "Duyệt người", "Sửa ràng buộc"], [
        ["[1] Playwright Test Agents", H, N, H, Y, H, N, H],
        ["[2] Codegen/locator", N, N, N, Y, H, H, N],
        ["[3] RBTG Survey", H, H, H, N, H, H, N],
        ["[4] WebTestPilot", Y, N, Y, Y, Y, N, H],
        ["[5] LLM Scenario Tool", H, N, N, N, H, H, N],
        ["[6] Zero-Cost Self-Healing", N, N, H, Y, H, N, Y],
        ["[7] GHL (ISO 29119-4)", H, N, H, N, H, H, N],
        ["[8] Autonomous Repair Limits", H, N, H, H, N, H, Y],
        ["[9] ReqToCode", Y, N, Y, N, N, H, N],
        ["[10] Steward", N, N, N, Y, H, N, H],
        ["[11] System TC (ChatGPT)", Y, N, H, N, H, Y, N],
        ["[12] Web Element Relocalization", N, N, N, Y, N, N, Y],
        ["Nhu cầu đề tài", Y, Y, Y, Y, Y, Y, Y],
    ], [4.2, 1.8, 1.7, 1.6, 1.5, 1.5, 1.7, 1.8], center_cols=(1, 2, 3, 4, 5, 6, 7))
    r.p("Cột phát hiện mơ hồ không có nguồn nào đạt \"Có\". [3] chỉ nêu đây là thách thức mở; [4] giả định đặc tả tự chứa và "
        "đầy đủ; [11] ghi nhận LLM hiểu hành vi hệ thống hạn chế khi chỉ đọc SRS nhưng không có bộ phát hiện. Không nguồn nào "
        "phủ đủ tám tiêu chí. Mỗi nguồn mạnh ở vài mảnh, nên khoảng trống mà đề tài nhắm tới là sự tích hợp có kiểm soát: một "
        "pipeline bán tự động có truy vết, có neo DOM, và ngăn self-healing làm đổi oracle khi chưa được duyệt.")
    r.p("Về RQ2, [10] báo cáo tỉ lệ hoàn thành tác vụ end-to-end 40% với GPT-4 khi có lọc DOM và screenshot, còn [8] cho thấy "
        "chiều ngược lại: khi agent bỏ qua kiểm chứng DOM, selector ảo giác gây lỗi dây chuyền. Chưa nguồn nào trong bộ khảo "
        "sát so sánh có kiểm soát cùng một tập test case khi có và không có ngữ cảnh accessibility tree. Thực nghiệm RQ2 của "
        "đồ án làm đúng phép so sánh này.")


def chapter3(r: Report):
    r.h1("Chương 3: PHÂN TÍCH, THIẾT KẾ VÀ XÂY DỰNG HỆ THỐNG", chapter=True)
    r.h2("3.1 Yêu cầu hệ thống")
    r.h3("3.1.1 Yêu cầu chức năng")
    r.bullets([
        "Nạp tài liệu DOCX, PDF, Markdown, text; trích yêu cầu có cấu trúc kèm trích đoạn nguồn.",
        "Gắn cờ yêu cầu thiếu, mơ hồ, mâu thuẫn và chờ người duyệt trước khi sinh test (AG-01).",
        "Sinh test case nhiều loại từ yêu cầu đã duyệt, cho người duyệt hoặc loại bỏ (AG-02).",
        "Sinh script Playwright Python từ test case, kiểm chứng locator trên trang thật.",
        "Chạy script, thu ảnh chụp, phân loại passed / failed / error / blocked.",
        "Khi test lỗi: phân mức rủi ro đề xuất sửa, tự chữa locator từ DOM sống, chờ người duyệt (AG-03..05).",
        "Dựng chuỗi truy vết, tính độ phủ yêu cầu, độ phủ thực thi, pass rate và artifact mồ côi.",
    ])
    r.h3("3.1.2 Yêu cầu phi chức năng")
    r.bullets([
        "Fail-closed: đầu ra sai schema, thiếu bằng chứng hay còn mơ hồ chưa xử lý thì dừng ở trạng thái chờ duyệt.",
        "Đổi backend LLM (cloud hoặc local) chỉ bằng cách đổi file cấu hình, không sửa agent.",
        "Kiểm thử offline: unit test không cần mạng, trình duyệt hay API key.",
        "Không lưu khoá API trong mã nguồn; .env và .env.cloud bị loại khỏi git.",
        "Tái lập được: mọi số liệu báo cáo chạy lại được bằng lệnh có ghi trong tài liệu.",
    ])

    r.h2("3.2 Kiến trúc tổng quan")
    r.p("Hệ thống gồm năm agent nối tiếp, mỗi agent dựa trên một dịch vụ tất định (Hình 3.1). Các agent không gọi nhau trực "
        "tiếp mà trao đổi qua các model Pydantic. Approval không phải công cụ mà agent tự gọi được: quyết định duyệt chỉ đến "
        "từ người dùng qua giao diện. Thiết kế này tránh việc agent tự hợp thức hoá đề xuất của chính nó.")
    r.figure(str(HERE / "fig_architecture.png"), "Kiến trúc tổng quan của hệ thống")
    r.p("Luồng xử lý đầy đủ đi qua các cổng duyệt như Hình 3.2. Khi test lỗi, Repair Agent tạo đề xuất; đề xuất được duyệt thì "
        "script mới được chạy lại; vượt ngân sách sửa thì chuyển sang trạng thái chờ người xem xét.")
    r.figure(str(HERE / "fig_pipeline.png"), "Luồng pipeline và các cổng phê duyệt")
    r.table("Các thư mục mã nguồn và vai trò", ["Thư mục / tệp", "Vai trò"], [
        ["src/models/", "Model Pydantic: Requirement, TestCase, PlaywrightPlan, ExecutionResult, RepairProposal, Approval"],
        ["src/agents/", "Năm agent: requirement, test_design, playwright_generation, execution, repair"],
        ["src/services/", "Dịch vụ tất định: llm_client, boundary_analysis, locator_policy, locator_healing, repair_policy, "
                          "script_template, traceability, document_loader"],
        ["src/services/browser_tasks.py", "Tác vụ trình duyệt dùng chung: snapshot, grounding, dò bước hỏng, chữa theo lô"],
        ["src/database/, src/repositories/", "Schema SQLAlchemy và repository (version, duyệt, truy vết, nhật ký)"],
        ["src/evaluation/", "Bộ đánh giá phát hiện mơ hồ: dataset, metrics, evaluator, design_metrics"],
        ["src/ui/", "Giao diện Streamlit (app.py) và logic duyệt thuần (review.py)"],
        ["apps/shoplab/", "Ứng dụng đích ShopLab, 14 mutation có nhãn, bộ 20 test"],
        ["run_*.py, evaluate_*.py", "CLI cho từng agent, pipeline đầy đủ và các thực nghiệm RQ1–RQ3"],
        ["tools/", "Công cụ thí nghiệm RQ4 và bộ dựng báo cáo"],
        ["datasets/reference/", "Ground truth RQ1 (15 tài liệu), catalog 50 target RQ2, nhiệm vụ RQ4"],
        ["tests/unit/", "Unit test offline"],
    ], [4, 12])

    r.h2("3.3 Requirement Analysis Agent")
    r.p("Agent nhận văn bản do bộ đọc tài liệu trích ra. Với PDF, bộ đọc chèn mốc [trang N] để mỗi yêu cầu truy được về trang "
        "gốc. LLM trả về danh sách StructuredRequirement theo schema ở Bảng 3.2. Hệ thống tự đánh ID REQ-001, REQ-002... theo thứ "
        "tự và bỏ qua ID do LLM đặt, vì ID của LLM không ổn định giữa các lần chạy. Tài liệu rỗng trả kết quả rỗng mà không gọi LLM.")
    r.table("Các trường của StructuredRequirement", ["Trường", "Ý nghĩa"], [
        ["id", "REQ-xxx do hệ thống gán, dùng làm gốc truy vết"],
        ["title, actor, action", "Tên yêu cầu, tác nhân thực hiện, hành động"],
        ["precondition, expected_outcome", "Điều kiện trước và kết quả mong đợi; để trống nếu tài liệu không nói"],
        ["constraints", "Ràng buộc (độ dài, khoảng giá trị...), đầu vào cho BVA"],
        ["source_excerpt", "Trích đoạn nguồn nguyên văn"],
        ["ambiguities", "Danh sách AmbiguityFinding: loại, mô tả, trích đoạn, gợi ý, mức độ"],
    ], [5, 11])
    r.p("Rubric mơ hồ có bảy loại (Bảng 3.3). Mâu thuẫn giữa nhiều yêu cầu được đặt ở global_ambiguities vì nó không thuộc "
        "riêng yêu cầu nào. Agent chỉ phát hiện, không tự làm đầy chỗ thiếu: yêu cầu bị gắn cờ dừng ở AG-01 cho tới khi người "
        "dùng làm rõ hoặc chấp nhận.")
    r.table("Rubric bảy loại mơ hồ", ["Loại", "Dấu hiệu", "Ví dụ"], [
        ["missing_actor", "Không rõ ai thực hiện", "\"Đơn hàng được hủy trong 24 giờ\""],
        ["missing_precondition", "Thiếu điều kiện trước cần thiết", "Đổi mật khẩu mà không nói phải đăng nhập"],
        ["vague_quantifier", "Lượng từ không đo được", "\"phản hồi nhanh\", \"một vài sản phẩm\""],
        ["missing_expected_outcome", "Không nói kết quả mong đợi", "\"Người dùng bấm Thanh toán\" rồi dừng"],
        ["underspecified_action", "Hành động thiếu tham số", "\"Nhập thông tin hợp lệ\""],
        ["conflict", "Hai yêu cầu mâu thuẫn", "\"mật khẩu 8–20 ký tự\" và \"tối thiểu 18 ký tự\""],
        ["other", "Mơ hồ khác", ""],
    ], [4.2, 5.3, 6.5])
    r.p("Một lỗi đáng ghi lại: khi bỏ phần mô tả schema khỏi prompt để tiết kiệm khoảng 35% token đầu vào, Gemma dồn mọi cờ "
        "mơ hồ vào global_ambiguities thay vì gắn theo từng yêu cầu. Micro-F1 rơi xuống 0,07. Thêm một dòng chỉ dẫn gắn cờ "
        "theo từng yêu cầu đưa micro-F1 lên 0,73. Bộ đánh giá định lượng là thứ làm lộ ra lỗi này; đọc mắt kết quả JSON thì "
        "khó thấy.")

    r.h2("3.4 Test Design Agent và Boundary Value Analysis")
    r.p("Agent chỉ nhận yêu cầu đã duyệt. LLM sinh test case các loại positive, negative, boundary, error_guessing và "
        "alternative_flow, mỗi test case mang ID REQ-xxx-TC-yy. Prompt cấm tạo expected result vượt quá yêu cầu đã duyệt.")
    r.p("LLM hay quên ca biên, nên hệ thống bổ sung module boundary_analysis tất định. Module nhận ràng buộc số tường minh "
        "(khoảng, tối thiểu, tối đa) viết bằng tiếng Việt, tiếng Anh hoặc ký hiệu. Từ mỗi ràng buộc, nó sinh ca biên hợp lệ và "
        "ca ngay ngoài biên, kèm trích dẫn ràng buộc nguồn. Ràng buộc không rõ thì không sinh (fail-closed). Với \"Mật khẩu "
        "8-20 ký tự\" và \"tối thiểu 18\", module sinh 6 ca biên và âm bản đúng nhãn.")

    r.h2("3.5 Playwright Generation Agent")
    r.h3("3.5.1 DSL action rút gọn")
    r.p("LLM không viết mã Python trực tiếp. Nó trả về một PlaywrightPlan gồm danh sách action trong một DSL nhỏ (Bảng 3.4); "
        "mã được render từ plan qua template Jinja2. Cách này giới hạn những gì LLM có thể làm, cho phép chính sách repair so "
        "sánh hai plan theo từng action, và loại bỏ lỗi cú pháp.")
    r.table("DSL action của PlaywrightPlan", ["Action", "Trường dùng", "Mã Playwright sinh ra"], [
        ["goto", "arg = URL", "page.goto(url)"],
        ["fill", "strategy, value, arg", "<locator>.fill(arg)"],
        ["click", "strategy, value, role_name", "<locator>.click()"],
        ["expect_url", "arg", "expect(page).to_have_url(arg)"],
        ["expect_visible", "strategy, value, role_name", "expect(<locator>).to_be_visible()"],
        ["expect_text", "strategy, value, arg", "expect(<locator>).to_contain_text(arg)"],
        ["(mọi locator)", "nth ≥ 0, tuỳ chọn", "<locator>.nth(nth) khi nhiều phần tử giống hệt nhau"],
    ], [3, 5, 8])
    r.p("Strategy là một trong role, label, placeholder, text, test_id, css. Với role, value là vai trò ARIA còn role_name là "
        "tên hiển thị. Prompt nhấn mạnh điểm này vì lỗi phổ biến nhất của LLM là đặt tên nút vào value. Trường nth được thêm "
        "sau khi RQ2 cho thấy một nhóm lỗi không giải được trong DSL cũ: trang có nhiều phần tử cùng vai trò và tên (sáu nút "
        "\"Add to cart\" của SauceDemo).")
    r.h3("3.5.2 Chính sách locator và grounding")
    r.p("Khi có nhiều locator ứng viên cho cùng phần tử, locator_policy chọn locator khớp đúng một phần tử, ưu tiên theo thứ tự "
        "role, label, placeholder, text, test_id, css như khuyến nghị của Playwright. Grounding đếm số phần tử khớp với mỗi "
        "locator trên trang thật và lưu thành GroundingRecord (ok khi khớp đúng 1).")
    r.p("Grounding tĩnh chụp DOM một lần nên bỏ sót phần tử chỉ có sau điều hướng, ví dụ danh sách sản phẩm chỉ hiện sau khi "
        "đăng nhập. Hàm ground_flow đi theo luồng: ở mỗi bước, nó chờ locator gắn vào DOM, đếm số phần tử khớp, rồi mới thực "
        "hiện goto, fill hoặc click để đẩy trang sang trạng thái kế. Assertion không làm đổi trạng thái. Test case đăng nhập "
        "SauceDemo đạt grounding 4/4, kể cả .inventory_list chỉ xuất hiện sau đăng nhập.")
    r.h3("3.5.3 Ngữ cảnh accessibility tree (nhánh DOM-aware)")
    r.p("Để trả lời RQ2, hàm plan nhận thêm tham số tuỳ chọn page_context. Khi có, prompt được nối thêm aria snapshot của "
        "trang đích lúc vừa mở, cắt tối đa 6.000 ký tự, kèm chỉ dẫn: với phần tử có trong snapshot thì dùng đúng role và tên "
        "như snapshot; phần tử xuất hiện sau điều hướng thì suy luận như thường. Khi không truyền page_context, prompt giữ "
        "nguyên như baseline. Unit test kiểm tra cả hai trường hợp.")
    r.h3("3.5.4 Vòng lặp grounding")
    r.p("Snapshot cho LLM biết phần tử nào tồn tại nhưng không cho biết locator nó viết ra có khớp đúng một phần tử hay "
        "không. Vòng lặp grounding ghép hai cơ chế: sau khi LLM sinh plan, ground_flow chạy plan trên trang thật và đếm số "
        "phần tử khớp ở từng bước. Nếu có locator khớp 0 hoặc nhiều hơn 1, hàm grounding_feedback dựng phản hồi gồm action "
        "lỗi, số phần tử khớp và aria snapshot của trang tại đúng bước đó (kể cả trang sau điều hướng), rồi gửi lại cho LLM "
        "để sinh lại đúng một lần. Script cuối cùng vẫn chỉ được chạy một lần.")

    r.h2("3.6 Execution Agent")
    r.p("Execution Agent ghi script ra tệp, chạy trong tiến trình con với timeout 120 giây và thu stdout, stderr, ảnh chụp. "
        "Trạng thái được phân loại: passed khi tiến trình thoát mã 0 và in PASSED; failed khi stderr chứa AssertionError, "
        "TimeoutError hoặc expect(; blocked khi quá timeout; các trường hợp còn lại là error. Agent không sửa test.")
    r.p("Template script cấu hình các tham số chạy, áp dụng như nhau cho mọi script. Timeout của expect là 10 giây vì một "
        "số trang demo cần khoảng 5 giây mới hiện kết quả, đúng bằng mặc định 5 giây của Playwright. Timeout thao tác "
        "(click, fill) là 10 giây thay cho mặc định 30 giây, để script hỏng theo chủ đích trong RQ3 không chiếm quá nhiều "
        "thời gian; timeout điều hướng giữ 30 giây. Chromium chạy với cờ --disable-http2, lý do trình bày ở mục 4.4.4.")

    r.h2("3.7 Repair Agent và chính sách sửa có ràng buộc")
    r.p("Repair Agent nhận test lỗi kèm bằng chứng và tạo đề xuất sửa dạng plan mới. Mức rủi ro do repair_policy quyết định "
        "bằng cách so sánh plan cũ và plan mới theo từng action, không để LLM tự khai (Bảng 3.5). Mọi đề xuất đều có "
        "requires_approval = true; agent không tự áp dụng. Ngân sách sửa mặc định là 2 lần cho một chuỗi thực thi; vượt ngân "
        "sách thì chuyển BLOCKED_FOR_REVIEW.")
    r.table("Phân mức rủi ro của đề xuất sửa", ["Mức", "Thay đổi phát hiện được", "Xử lý"], [
        ["Low", "Chỉ đổi locator", "Đề xuất, người duyệt đơn giản"],
        ["Medium", "Đổi dữ liệu nhập, URL điều hướng, thêm bước", "Bắt buộc xem bằng chứng và duyệt"],
        ["High", "Đổi assertion, đổi loại action, bỏ bước", "Bắt buộc duyệt, gắn cờ thay đổi ngữ nghĩa"],
        ["Prohibited", "Sửa mã nguồn website, xoá bằng chứng, lặp vô hạn", "Từ chối"],
    ], [2.6, 7.4, 6])
    r.p("Self-healing locator là đường sửa không cần LLM. Khi locator gãy, hệ thống quét DOM sống lấy các ứng viên, so khớp "
        "mờ với ý định của locator cũ, đếm số phần tử khớp, rồi chọn ứng viên duy nhất theo locator_policy. Không có ứng viên "
        "duy nhất thì trả về rỗng. Trong demo, locator cố ý sai placeholder:Usernamex được chữa thành role:textbox[Username] "
        "với rủi ro Low, và vẫn chờ người duyệt.")
    r.p("Lần chạy RQ3 đầu cho thấy một giới hạn: self-healing chỉ chữa locator hỏng đầu tiên mỗi lượt, nên mutation đổi id "
        "của cả ba ô đăng nhập cần ba lượt và vượt ngân sách 2. Hàm heal_all chữa theo lô: dò bước hỏng, chữa, chạy lại plan "
        "đã chữa tới bước hỏng kế tiếp, lặp tới khi không còn locator chữa được; mọi chỗ chữa gộp thành một đề xuất. Mức "
        "rủi ro vẫn do policy quyết định trên toàn bộ thay đổi. Tập ứng viên của thao tác click được mở rộng thêm "
        "input[type=submit] và input[type=button], vì nút Login của SauceDemo là một thẻ input.")
    r.p("RepairProposal mang theo plan mới (new_plan) bên cạnh diff. Đề xuất chỉ được áp dụng sau khi người duyệt đồng ý; "
        "khi đó plan mới được lưu thành version script mới và chạy lại.")

    r.h2("3.8 Truy vết, cổng duyệt và giao diện")
    r.p("Traceability service dựng chuỗi Requirement → TestCase → Script → Execution → Repair từ các ID, tính độ phủ yêu cầu "
        "(tỉ lệ yêu cầu có ít nhất một test case), độ phủ thực thi, pass rate và liệt kê artifact mồ côi. Bảng 3.6 liệt kê các "
        "cổng duyệt.")
    r.table("Các cổng phê duyệt", ["Cổng", "Khi nào", "Trạng thái hiện thực"], [
        ["AG-01", "Yêu cầu bị gắn cờ mơ hồ", "Có trên giao diện Streamlit"],
        ["AG-02", "Trước khi sinh script từ test case", "Có trên giao diện Streamlit"],
        ["AG-03", "Đề xuất sửa đổi phần tử mà assertion kiểm", "Có trên giao diện Streamlit"],
        ["AG-04", "Đề xuất đổi giá trị kỳ vọng (oracle)", "Có trên giao diện Streamlit"],
        ["AG-05", "Đề xuất bỏ, thêm hoặc đổi bước kiểm thử", "Có trên giao diện Streamlit"],
    ], [2, 7, 7], center_cols=(0,))
    storage_section(r)
    ui_section(r)

    r.h2("3.11 Lớp trừu tượng LLM")
    r.p("Mọi agent gọi LLM qua giao diện LLMClient với một phương thức structured_completion(system, user, schema). Có hai hiện "
        "thực: OpenAILLMClient gọi mọi endpoint tương thích OpenAI (OpenAI, DeepSeek, LM Studio, Ollama) bằng base_url, và "
        "MockLLMClient trả đối tượng dựng sẵn cho test. Tham số --profile cloud đọc .env.cloud thay cho .env, nên đổi giữa "
        "Gemma chạy local và DeepSeek chỉ cần một cờ dòng lệnh. LM Studio từ chối response_format kiểu json_object, nên hệ "
        "thống dùng json_schema, kiểu được cả ba backend chấp nhận.")

    r.h2("3.12 Kiểm thử phần mềm")
    r.p(f"Bộ unit test gồm {UNIT_TESTS} test chạy offline bằng MockLLMClient và hàm đếm giả, không cần mạng hay trình duyệt. "
        "Các phần tất định được test kỹ: repair policy, traceability, self-healing, boundary analysis, locator policy, bộ đánh "
        "giá RQ1, repository trên SQLite in-memory, hành vi các mutation của ShopLab (qua Flask test client), ánh xạ cổng "
        "duyệt, chấm điểm SUS và logic duyệt của giao diện. Phần cần trình duyệt thật (grounding đa bước, chạy script, chữa "
        "theo lô, luồng giao diện) được kiểm chứng bằng chạy thật: oracle RQ2, bước --check của RQ3 và kịch bản Playwright "
        "đi hết luồng giao diện.")


def chapter4(r: Report, cat):
    r.h1("Chương 4: MÔI TRƯỜNG THỰC NGHIỆM VÀ PHƯƠNG PHÁP ĐÁNH GIÁ", chapter=True)
    r.h2("4.1 Môi trường thực nghiệm")
    r.p("Thực nghiệm chạy trên một máy Windows 11, Python 3.14, Playwright 1.62 với Chromium headless. Backend cloud là "
        "DeepSeek (model deepseek-chat) qua API tương thích OpenAI. Backend local là google/gemma-4-12b chạy trên LM Studio với "
        "GPU 8 GB. Các website đích đều là site demo công khai dựng cho mục đích luyện kiểm thử tự động.")

    shoplab_section(r)

    r.h2("4.3 Đánh giá RQ1: yêu cầu có cấu trúc")
    r.h3("4.3.1 Dataset")
    r.p("Ground truth được tạo bằng cách chèn khuyết tật có chủ đích vào tài liệu yêu cầu. Tập ban đầu gồm 5 tài liệu (auth, "
        "booking, cart, profile, search) với 15 yêu cầu. Tập mở rộng thêm 10 tài liệu thuộc các nghiệp vụ khác (thư viện, "
        "chuyển tiền, đăng ký học phần, kho, đặt bàn, ví điện tử, nghỉ phép, mạng xã hội nội bộ, mã giảm giá, hỗ trợ khách "
        "hàng), tổng 15 tài liệu và 47 yêu cầu. Khoảng một nửa yêu cầu được chèn một khuyết tật thuộc rubric, hai tài liệu "
        "có mâu thuẫn toàn cục; các yêu cầu sạch thường có ràng buộc số để kiểm khả năng sinh ca biên. Mỗi tài liệu có thêm "
        "danh sách điều kiện kiểm thử gold (tổng 83 điều kiện: 41 positive, 18 negative, 24 biên), là các hành vi một kiểm "
        "thử viên rút được từ văn bản đúng như đang viết. 10 tài liệu mới và toàn bộ điều kiện gold do Claude soạn nháp, "
        "nhóm cần review trước khi dùng làm ground truth chính thức.")
    r.h3("4.3.2 Phát hiện mơ hồ")
    r.p("Đơn vị so khớp là cặp (yêu cầu, loại mơ hồ); yêu cầu dự đoán được căn với yêu cầu gốc theo thứ tự use case; mâu "
        "thuẫn được tính như một loại ở mức toàn cục. Yêu cầu sạch dùng để đo over-flag. Chạy: python evaluate_agent.py "
        "(thêm --profile cloud cho DeepSeek).")
    r.h3("4.3.3 Chất lượng test case: hai nhánh")
    r.p("Đây là phần trả lời trực tiếp câu hỏi RQ1. Nhánh direct đưa cả tài liệu cho LLM và yêu cầu sinh test case. Nhánh "
        "pipeline chạy Requirement Agent, duyệt mô phỏng mọi yêu cầu, rồi chạy Test Design Agent và BVA trên từng yêu cầu. "
        "Hai nhánh dùng cùng model và cùng các quy tắc chống bịa trong prompt. Một LLM giám khảo nhận tài liệu, danh sách "
        "điều kiện gold và các test case, rồi với mỗi test case trả về các điều kiện nó thực sự kiểm và việc kết quả mong "
        "đợi của nó có căn cứ trong tài liệu hay không. Từ đó tính độ phủ (tỉ lệ điều kiện gold có ít nhất một test kiểm), "
        "tỉ lệ test không căn cứ và tỉ lệ test truy vết được về một yêu cầu có ID. Một mẫu ngẫu nhiên 20% phán quyết được "
        "xuất ra CSV để người chấm lại và tính hệ số kappa (python evaluate_rq1.py --kappa). Lệnh chạy: python "
        "evaluate_rq1.py --profile cloud.")

    r.h2("4.4 Đánh giá RQ2: ngữ cảnh DOM và first-run pass rate")
    r.h3("4.4.1 Catalog target")
    sites = Counter(t["site"] for t in cat)
    types = Counter(t["type"] for t in cat)
    r.p(f"Catalog datasets/reference/rq2_targets.json gồm {len(cat)} target trên {len(sites)} website (Bảng 4.2), "
        f"{types['positive']} positive và {types['negative']} negative. Mỗi target là một test case viết bằng tiếng Việt (các "
        "bước và expected_result) trên một URL thật. expected_result mô tả kết quả thật của trang để LLM có căn cứ sinh "
        "assertion; RQ2 đo khả năng chuyển test case đã duyệt thành mã, không đo khả năng đoán oracle. Danh sách đầy đủ ở "
        "Phụ lục A.")
    rows = []
    for s, n in sorted(sites.items(), key=lambda kv: -kv[1]):
        sample = next(t for t in cat if t["site"] == s)
        kinds = {
            "saucedemo": "Đăng nhập, giỏ hàng, checkout, chi tiết sản phẩm, đăng xuất",
            "the-internet": "Đăng nhập, JS alert/confirm, nội dung động, redirect, status code",
            "uitestingplayground": "Form đăng nhập, đổi tên nút, đếm click, điều hướng",
            "practicetestautomation": "Đăng nhập, đăng xuất, thêm dòng động",
            "demoqa": "Form text box, radio button, nút bấm",
            "orangehrm": "Đăng nhập và validation",
        }[s]
        rows.append([SITE_NAME[s], str(n), kinds])
    r.table("Các website trong catalog RQ2", ["Website", "Số target", "Loại kịch bản"], rows, [4.6, 2, 9.4], center_cols=(1,))
    r.h3("4.4.2 Oracle viết tay: kiểm chứng target khả thi")
    r.p("Một target có thể thất bại vì lý do không liên quan tới LLM: expected_result viết sai, site thay đổi, hoặc kịch bản "
        "cần thao tác ngoài DSL (chọn dropdown, hover, nhấn phím). Để loại các nguồn nhiễu này, mỗi target có một trường "
        "oracle là plan viết tay bằng đúng DSL mà LLM được dùng. Lệnh python run_rq2.py --context oracle chạy các plan này "
        "qua cùng template và Execution Agent.")
    r.p("Trong lúc xây catalog, oracle đã loại nhiều ứng viên. TodoMVC bị loại vì thêm việc cần nhấn Enter. Các trang AJAX và "
        "client delay của UI Testing Playground bị loại vì kết quả hiện sau khoảng 15 giây. Dropdown, hover, key press và "
        "double click bị loại vì DSL không có thao tác tương ứng. Oracle cũng bắt được lỗi trong catalog cũ: OrangeHRM hiển thị "
        "chữ \"Required\" dưới cả hai ô khi để trống cả hai, nên assertion text khớp hai phần tử và vi phạm strict mode. Target "
        "này được sửa thành bỏ trống riêng ô mật khẩu.")
    r.h3("4.4.3 Các nhánh so sánh")
    r.p("Cả hai nhánh dùng cùng model, cùng system prompt, cùng 50 test case và cùng template thực thi. Chỉ khác ở ngữ cảnh "
        "đưa thêm vào prompt (Hình 4.1):")
    r.bullets([
        "none (baseline): prompt gồm URL đích, tiêu đề, loại, các bước và expected_result của test case.",
        "aria (DOM-aware): prompt như baseline, thêm aria snapshot của trang đích chụp ngay trước khi gọi LLM.",
        "aria_loop (bản 2): như aria, thêm vòng lặp grounding của mục 3.5.4, tức sinh lại một lần nếu có locator khớp khác 1.",
    ])
    r.p("Lần đo đầu (mục 5.2.1–5.2.2) chỉ có hai nhánh none và aria, chạy một lần. Bản 2 chạy cả ba nhánh, lặp ba lần để "
        "ước lượng độ dao động giữa các lần sinh của LLM, với template mới (timeout thao tác 10 giây, trường nth). Với "
        "template này, oracle đạt 49/50: PT-04 (bấm Log out trên Practice Test Automation) vượt timeout 10 giây vì trang "
        "điều hướng chậm. Giới hạn này áp dụng như nhau cho mọi nhánh.")
    r.figure(str(HERE / "fig_rq2_protocol.png"), "Quy trình đo RQ2 trên hai nhánh")
    r.p("Script sinh ra được chạy đúng một lần, không qua Repair Agent hay sửa tay. Plan không hợp lệ (LLM trả JSON hỏng) được "
        "tính là error. Mỗi kết quả lưu kèm dòng lỗi cuối của stderr để phân loại nguyên nhân. Lệnh chạy: python run_rq2.py "
        "--profile cloud --context both --out rq2_results_both.json.")
    r.h3("4.4.4 Kiểm soát nhiễu hạ tầng")
    r.p("Khi chạy oracle lần đầu, mọi target của The Internet đều fail ở page.goto do timeout 30 giây, trong khi curl tải đúng "
        "các tài nguyên đó dưới 1 giây. Theo dõi request cho thấy Chromium kẹt khi tải tệp js và png tĩnh qua HTTP/2 từ mạng "
        "của nhóm. Thêm cờ --disable-http2 thì cả ba lần thử đều tải xong trong 3,3 giây. Lỗi này cũng có thể đã làm giảm kết "
        "quả của lần đo thử ngày 15/09 (8/19 target đạt), trong đó The Internet chỉ đạt 3/7.")
    r.p("Sau khi tắt HTTP/2, The Internet vẫn thỉnh thoảng treo khi tải trang. Đây là lỗi hạ tầng, không phải lỗi script, nên "
        "khi một script fail với thông điệp Page.goto: Timeout, hệ thống chạy lại đúng script đó một lần và đánh dấu "
        "infra_retry trong JSON. Script không đổi nên phép đo first-run vẫn giữ nguyên ý nghĩa: LLM chỉ có một cơ hội sinh mã.")

    rq3_method(r, RQ3_COUNTS)
    rq4_method(r)


def chapter5(r: Report, res, cat, v1):
    r.h1("Chương 5: KẾT QUẢ THỰC NGHIỆM VÀ ĐÁNH GIÁ", chapter=True)
    r.h2("5.1 RQ1: yêu cầu có cấu trúc")
    r.h3("5.1.1 Phát hiện mơ hồ trên tập ban đầu")
    r.p("Bảng 5.1 là kết quả trên 15 yêu cầu sau khi siết prompt. Dòng Gemma \"chưa nắn prompt\" là cấu hình trước khi thêm "
        "chỉ dẫn gắn cờ theo từng yêu cầu (mục 3.3).")
    r.table("Kết quả phát hiện mơ hồ trên tập 15 yêu cầu", ["Model", "Micro-P", "Micro-R", "Micro-F1", "Macro-F1", "Over-flag"], [
        ["deepseek-chat (cloud)", "0,50", "1,00", "0,67", "0,71", "0,20"],
        ["gemma-4-12b, chưa nắn prompt", "0,06", "0,08", "0,07", "0,17", "0,00"],
        ["gemma-4-12b (local)", "0,71", "0,83", "0,77", "0,69", "0,00"],
    ], [5.5, 2.1, 2.1, 2.1, 2.1, 2.1], center_cols=(1, 2, 3, 4, 5))
    r.p("Siết prompt (chỉ gắn cờ khi thiếu hoặc mơ hồ rõ ràng, ưu tiên không gắn khi phân vân ở precondition và outcome) giảm "
        "over-flag ở cả hai backend (Bảng 5.2). Đổi lại, recall của Gemma giảm từ 1,00 xuống 0,83, chủ yếu do bỏ sót "
        "missing_precondition; DeepSeek giữ recall 1,00.")
    r.table("Ảnh hưởng của việc siết prompt (trước → sau)", ["Model", "Micro-F1", "Over-flag", "Recall"], [
        ["gemma-4-12b (local)", "0,73 → 0,77", "0,40 → 0,00", "1,00 → 0,83"],
        ["deepseek-chat (cloud)", "0,51 → 0,67", "0,80 → 0,20", "1,00 → 1,00"],
    ], [5.5, 3.5, 3.5, 3.5], center_cols=(1, 2, 3))
    r.table("F1 theo từng loại mơ hồ", ["Loại", "deepseek-chat", "gemma-4-12b"], [
        ["conflict", "1,00", "1,00"], ["missing_actor", "0,86", "1,00"], ["missing_expected_outcome", "0,60", "0,67"],
        ["missing_precondition", "0,50", "0,00"], ["underspecified_action", "0,29", "0,50"], ["vague_quantifier", "1,00", "1,00"],
    ], [7, 4.5, 4.5], center_cols=(1, 2))
    r.p("Dưới cùng một prompt, Gemma có precision và micro-F1 cao hơn, DeepSeek bắt đủ hơn (recall và macro-F1 cao hơn). "
        "missing_precondition là loại khó nhất với cả hai. Nhãn gold của loại này cũng chủ quan nhất: mỗi yêu cầu chỉ có một "
        "khuyết tật được chèn, trong khi một người đọc khác có thể thấy thêm precondition còn thiếu. Với 15 yêu cầu, các con số "
        "này chỉ là sơ bộ; một yêu cầu đổi nhãn có thể làm F1 của một loại thay đổi vài chục điểm phần trăm.")
    rq1_ambiguity(r, DATA["amb47"], notes_ext)
    rq1_testgen(r, DATA["rq1_v1"], DATA["rq1_v2"], notes_ext)

    rq2_section(r, res, cat)
    rq2_v1_section(r, v1, res)
    r.h3("5.2.2 Thảo luận về lần đo đầu")
    for para in RQ2_DISCUSSION:
        r.p(para)
    rq2_v2_section(r, DATA["rq2_runs"], notes_ext, DATA["rq2_vac"])

    rq3_results(r, DATA["rq3"], DATA["rq3_first"], notes_ext)
    rq4_results(r, DATA["rq4"], notes_ext)
    r.h2("5.5 Mối đe dọa tính hợp lệ")
    r.bullets(notes_ext.THREATS)


def rq2_section(r: Report, res, cat):
    none, aria = res["conditions"]["none"], res["conditions"]["aria"]
    paired = res["paired"]
    n = none["n"]
    p_val = mcnemar_exact_p(paired["aria_only"], paired["none_only"])
    r.h2("5.2 RQ2: ngữ cảnh DOM và first-run pass rate")
    r.p(f"Oracle viết tay đạt {ORACLE_PASS}/{n} target, nên catalog khả thi với DSL hiện có và đóng vai cận trên của phép đo. "
        f"Với model {res['model']}, nhánh baseline đạt {none['passed']}/{n} ({pct(none['passed'], n)}), nhánh có aria snapshot "
        f"đạt {aria['passed']}/{n} ({pct(aria['passed'], n)}).")
    rows = []
    for s in sorted(none["by_site"], key=lambda k: -none["by_site"][k]["n"]):
        a, b = none["by_site"][s], aria["by_site"][s]
        rows.append([SITE_NAME[s], str(a["n"]), f"{a['passed']} ({pct(a['passed'], a['n'])})",
                     f"{b['passed']} ({pct(b['passed'], b['n'])})"])
    rows.append(["Tổng", str(n), f"{none['passed']} ({pct(none['passed'], n)})", f"{aria['passed']} ({pct(aria['passed'], n)})"])
    r.table("First-run pass rate theo website, hai nhánh", ["Website", "Số target", "Baseline (none)", "DOM-aware (aria)"],
            rows, [5.2, 2.4, 4.2, 4.2], center_cols=(1, 2, 3))
    r.figure(str(HERE / "fig_rq2_sites.png"), "Tỉ lệ đạt lần đầu theo website, baseline và DOM-aware", width_cm=15)

    r.p("Hai nhánh chạy trên cùng target nên có thể so sánh từng cặp (bảng dưới).")
    r.table("Bảng chéo kết quả theo cặp target", ["", "aria đạt", "aria trượt"], [
        ["none đạt", str(paired["both_pass"]), str(paired["none_only"])],
        ["none trượt", str(paired["aria_only"]), str(paired["both_fail"])],
    ], [5, 5.5, 5.5], center_cols=(1, 2))
    r.p(f"Có {paired['aria_only']} target chỉ nhánh aria đạt và {paired['none_only']} target chỉ baseline đạt. Kiểm định McNemar "
        f"chính xác trên {paired['aria_only'] + paired['none_only']} cặp bất đồng cho p = {p_val:.3f}".replace('.', ',')
        + (", nhỏ hơn 0,05: khác biệt có ý nghĩa thống kê ở cỡ mẫu này." if p_val < 0.05 else
           ", không nhỏ hơn 0,05, nên với cỡ mẫu 50 chưa đủ bằng chứng thống kê để khẳng định khác biệt."))

    cn = Counter(filter(None, (cause(c) for c in none["cases"])))
    ca = Counter(filter(None, (cause(c) for c in aria["cases"])))
    keys = sorted(set(cn) | set(ca), key=lambda k: -(cn[k] + ca[k]))
    r.table("Nguyên nhân thất bại theo nhánh", ["Nguyên nhân (từ dòng lỗi cuối của stderr)", "Baseline", "DOM-aware"],
            [[k, str(cn[k]), str(ca[k])] for k in keys] + [["Tổng số target trượt", str(sum(cn.values())), str(sum(ca.values()))]],
            [9, 3.5, 3.5], center_cols=(1, 2))
    retries = sum(c.get("infra_retry", False) for c in none["cases"] + aria["cases"])
    r.p(f"Trong {2 * n} lần chạy có {retries} lần cần chạy lại do Page.goto: Timeout (mục 4.4.4).")


def rq2_v1_section(r: Report, v1, v2):
    """Lần đo thứ nhất (prompt thiếu đặc tả goto) — giữ lại làm bằng chứng."""
    r.h3("5.2.1 Lần đo thứ nhất: một lỗi đặc tả prompt bị che khuất")
    a1, b1 = v1["conditions"]["none"], v1["conditions"]["aria"]
    g1 = sum("invalid URL" in c["error"] for c in a1["cases"])
    g2 = sum("invalid URL" in c["error"] for c in b1["cases"])
    n = a1["n"]
    r.p(f"Lần đo đầu tiên cho baseline {a1['passed']}/{n} ({pct(a1['passed'], n)}) và aria {b1['passed']}/{n} "
        f"({pct(b1['passed'], n)}). Đọc từng script lỗi thì thấy {g1} script baseline bắt đầu bằng page.goto(\"\"), trong khi "
        f"nhánh aria chỉ có {g2}. Gọi lại LLM trên ba target cho thấy DeepSeek đều đặn đặt URL vào trường value thay vì arg. "
        "Nguyên nhân nằm ở system prompt: phần mô tả trường có nói fill dùng arg cho giá trị nhập và expect_url dùng arg cho URL, "
        "nhưng bỏ sót goto. Khi prompt dài thêm vì có aria snapshot, model lại điền đúng.")
    r.p("Nếu dừng ở lần đo này, khoảng cách giữa hai nhánh sẽ bị quy hết cho ngữ cảnh DOM, trong khi phần lớn nó đến từ một "
        "lỗi đặc tả không liên quan tới locator. Nhóm thêm một dòng \"goto: arg = URL\" vào prompt (áp dụng cho cả hai nhánh) "
        "rồi đo lại toàn bộ. Các kết quả đầu mục 5.2 là của lần đo thứ hai; bảng dưới đặt hai lần đo cạnh nhau.")
    a2, b2 = v2["conditions"]["none"], v2["conditions"]["aria"]
    g1b = sum("invalid URL" in c["error"] for c in a2["cases"])
    g2b = sum("invalid URL" in c["error"] for c in b2["cases"])
    r.table("Hai lần đo RQ2 trước và sau khi sửa đặc tả goto", ["Lần đo", "Baseline đạt", "Aria đạt", "goto rỗng (baseline / aria)"], [
        ["1 (prompt thiếu đặc tả goto)", f"{a1['passed']}/{n}", f"{b1['passed']}/{n}", f"{g1} / {g2}"],
        ["2 (đã sửa prompt)", f"{a2['passed']}/{n}", f"{b2['passed']}/{n}", f"{g1b} / {g2b}"],
    ], [6, 3, 3, 4], center_cols=(1, 2, 3))


def cause(c):
    """Nhóm nguyên nhân fail từ dòng lỗi cuối của stderr."""
    e = c["error"]
    if c["status"] == "passed":
        return None
    if e.startswith("plan:"):
        return "Plan không hợp lệ (JSON/schema)"
    if "invalid URL" in e:
        return "Bước goto thiếu URL (đặt sai trường)"
    if "strict mode violation" in e or e.startswith("1) <"):
        return "Locator khớp nhiều phần tử (strict mode)"
    if "Page.goto: Timeout" in e:
        return "Site không tải được (hạ tầng)"
    if "to_have_url" in e or "Page URL expected" in e:
        return "Assertion URL sai"
    if "Locator expected" in e or "AssertionError" in e or "to_contain_text" in e or "to_be_visible" in e:
        return "Assertion sai hoặc không thấy phần tử kiểm chứng"
    if "Timeout" in e or "waiting for" in e:
        return "Không tìm thấy phần tử để thao tác (timeout)"
    return "Lỗi khác"


def chapter6(r: Report):
    r.h1("Chương 6: KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", chapter=True)
    r.h2("6.1 Kết quả đạt được")
    r.bullets(notes_ext.CONCLUSIONS)
    r.h2("6.2 Hạn chế")
    r.bullets(notes_ext.LIMITS)
    r.h2("6.3 Hướng phát triển")
    r.bullets(notes_ext.FUTURE)


REFERENCES = [
    "Microsoft. Playwright Test Agents. https://playwright.dev/docs/test-agents, 2025.",
    "Microsoft. Playwright: Codegen và Best Practices. https://playwright.dev/docs/best-practices, 2026.",
    "Z. Yang, R. Huang, C. Cui, N. Niu, D. Towey. Requirements-Based Test Generation: A Comprehensive Survey. arXiv:2505.02015, 2025.",
    "X. Teoh, Y. Lin, D.-M. Nguyen, R. Ren, W. Zhang, J. S. Dong. WebTestPilot. FSE 2026, arXiv:2602.11724, DOI 10.1145/3797115.",
    "A. M. Sami, Z. Rasheed, M. Waseem, Z. Zhang, T. Herda, P. Abrahamsson. A Tool for Test Case Scenarios Generation Using "
    "Large Language Models. arXiv:2406.07021, 2024.",
    "R. N. Joseph. Beyond LLM-Based Test Automation: Zero-Cost Self-Healing via DOM Accessibility Tree. arXiv:2603.20358, 2026.",
    "S. Masuda, S. Kouzawa, K. Sezai, H. Suhara, Y. Hiruta, K. Kudou. Generating High-Level Test Cases from Requirements using "
    "LLM: An Industry Study. arXiv:2510.03641, 2025.",
    "H. Lee. Practical Limits of Autonomous Test Repair: A Multi-Agent Case Study. arXiv:2605.01471, 2026.",
    "T. Schlathölter. ReqToCode: Embedding Requirements Traceability as a Structural Property of the Codebase. "
    "arXiv:2603.13999, 2026.",
    "B. Tang, K. G. Shin. Steward: Natural Language Web Automation. arXiv:2409.15441, 2024.",
    "S. Bhatia, T. Gandhi, D. Kumar, P. Jalote. System Test Case Design from Requirements Specifications: Insights and "
    "Challenges of Using ChatGPT. arXiv:2412.03693, 2024.",
    "A. Kluge, A. Stocco. Web Element Relocalization in Evolving Web Applications. arXiv:2505.16424, 2025.",
]


def references_and_appendix(r: Report, cat):
    r.h1("TÀI LIỆU THAM KHẢO")
    for i, ref in enumerate(REFERENCES, 1):
        r.p(f"[{i}] {ref}", align=WD_ALIGN_PARAGRAPH.LEFT)
    r.h1("PHỤ LỤC A: DANH SÁCH 50 TARGET RQ2")
    r.chapter = "A"
    r.n_table = 0
    rows = [[t["id"], SITE_NAME[t["site"]], t["type"], "; ".join(t["steps"]), t["expected_result"]] for t in cat]
    r.table("Catalog target RQ2 (datasets/reference/rq2_targets.json)", ["ID", "Website", "Loại", "Các bước", "Kết quả mong đợi"],
            rows, [1.5, 2.6, 1.8, 5.3, 4.8], center_cols=(0, 2))

    from apps.shoplab.suite import SUITE
    r.h1("PHỤ LỤC B: BỘ 20 TEST SHOPLAB (RQ3)")
    r.chapter, r.n_table = "B", 0
    r.table("Bộ test ShopLab (apps/shoplab/suite.py)", ["ID", "Loại", "Tiêu đề", "Các bước", "Kết quả mong đợi"],
            [[tid, kind, title, "; ".join(steps), exp] for tid, kind, title, steps, exp, _ in SUITE],
            [1.3, 1.8, 3.4, 5.6, 3.9], center_cols=(0, 1))

    gold = json.loads((ROOT / "datasets/reference/ambiguity_eval/test_conditions.json").read_text(encoding="utf-8"))["cases"]
    r.h1("PHỤ LỤC C: ĐIỀU KIỆN KIỂM THỬ GOLD (RQ1)")
    r.chapter, r.n_table = "C", 0
    r.table("83 điều kiện kiểm thử gold (datasets/reference/ambiguity_eval/test_conditions.json)",
            ["Tài liệu", "ID", "UC", "Loại", "Điều kiện"],
            [[doc.removeprefix("cases/").removesuffix(".md"), g["id"], str(g["uc"]), g["kind"], g["condition"]]
             for doc, gs in gold.items() for g in gs], [2.3, 1.2, 1.1, 2, 9.4], center_cols=(1, 2, 3))

    r.h1("PHỤ LỤC D: HƯỚNG DẪN CHẠY LẠI")
    r.chapter, r.n_table = "D", 0
    r.p("Cài đặt: pip install -r requirements.txt rồi python -m playwright install chromium. Cấu hình LLM trong .env (local) "
        "và .env.cloud (DeepSeek), xem .env.example; hai file này không được commit. Mọi lệnh chạy ở thư mục gốc dự án.",
        align=WD_ALIGN_PARAGRAPH.LEFT)
    r.table("Lệnh tái lập các số liệu trong báo cáo", ["Mục", "Lệnh"], [
        ["Unit test", "python -m pytest tests/unit -q"],
        ["Giao diện", "streamlit run src/ui/app.py"],
        ["ShopLab", "python apps/shoplab/app.py --variant v0 --port 5100"],
        ["RQ1 phát hiện mơ hồ", "python evaluate_agent.py --profile cloud"],
        ["RQ1 chất lượng test case", "python evaluate_rq1.py --profile cloud --out rq1_results_v2.json"],
        ["RQ1 kappa", "python evaluate_rq1.py --kappa rq1_judge_sample_v2.csv"],
        ["RQ2 kiểm catalog", "python run_rq2.py --context oracle"],
        ["RQ2 lần đo đầu", "python run_rq2.py --profile cloud --context both --out rq2_results_both.json"],
        ["RQ2 bản 2", "python run_rq2.py --profile cloud --context all --out rq2_v2_r1.json (lặp r2, r3)"],
        ["RQ3 kiểm suite", "python run_rq3.py --check"],
        ["RQ3 ba nhánh", "python run_rq3.py --profile cloud --out rq3_results_v2.json"],
        ["RQ4", "python tools/rq4/rq4.py timer | export | check | analyze (xem rq4-protocol.md)"],
    ], [4.5, 11.5])


import notes_ext  # noqa: E402  (nhận xét viết sau khi có kết quả thật)
from chapters_ext import load as ext_load  # noqa: E402
from chapters_ext import rq3_method, rq4_method, shoplab_section, storage_section, ui_section  # noqa: E402
from chapters_results import rq1_ambiguity, rq1_testgen, rq2_v2_section, rq3_results, rq4_results  # noqa: E402
from rq2_notes import ORACLE_PASS, RQ2_DISCUSSION, UNIT_TESTS  # noqa: E402  (viết sau khi có kết quả thật)

RQ4_STATUS = notes_ext.RQ4_STATUS
DATA: dict = {}
RQ3_COUNTS: dict = {}


def _load_data() -> None:
    """Nạp mọi file kết quả đợt 2 (thiếu file nào thì mục tương ứng ghi 'chưa có kết quả')."""
    DATA["amb47"] = ext_load("datasets/processed/eval_deepseek-chat.json")
    DATA["rq1_v1"] = ext_load("rq1_results.json")
    DATA["rq1_v2"] = ext_load("rq1_results_v2.json")
    DATA["rq2_runs"] = [d for i in (1, 2, 3) if (d := ext_load(f"rq2_v2_r{i}.json"))]
    DATA["rq2_vac"] = [ext_load(f"rq2_v2_vac_r{i}.json") for i in range(1, len(DATA["rq2_runs"]) + 1)]
    rerun = ext_load("rq3_results_v2.json")
    DATA["rq3"], DATA["rq3_first"] = (rerun, ext_load("rq3_results.json")) if rerun else (ext_load("rq3_results.json"), None)
    DATA["rq4"] = ext_load("artifacts/rq4/rq4_results.json")
    counts = Counter(x["variant"] for x in DATA["rq3"]["cases"] if x["arm"] == "constrained")
    RQ3_COUNTS.update(counts)


def build(results_path: str, v1_path: str, out: str):
    _load_data()
    res = json.loads(Path(results_path).read_text(encoding="utf-8"))
    v1 = json.loads(Path(v1_path).read_text(encoding="utf-8"))
    cat = json.loads((ROOT / "datasets/reference/rq2_targets.json").read_text(encoding="utf-8"))["targets"]
    r = Report(TEMPLATE, "Tp. Hồ Chí Minh")
    front_matter(r)
    chapter1(r)
    chapter2(r)
    chapter3(r)
    chapter4(r, cat)
    chapter5(r, res, cat, v1)
    chapter6(r)
    references_and_appendix(r, cat)
    r.update_fields_on_open()
    r.save(out)
    print(f"[✓] {out}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    build(sys.argv[1], sys.argv[2], sys.argv[3])
