"""Helper dựng báo cáo trên nền file mẫu (giữ style CapBang/CapHinh/CellText/CodeBlock/List Bullet)."""
from __future__ import annotations

import copy
import re

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

TEXT_WIDTH_DXA = 9071  # A4, lề trái 3cm + phải 2cm -> 16cm
_INLINE = re.compile(r"(`[^`]+`)")


class Report:
    def __init__(self, template: str, keep_until_text: str):
        self.doc = docx.Document(template)
        self.chapter = 0
        self.n_table = 0
        self.n_fig = 0
        self._truncate_after(keep_until_text)

    # ---------- khung ----------
    def _truncate_after(self, marker: str):
        body = self.doc.element.body
        anchor = next(p._p for p in self.doc.paragraphs if p.text.strip().startswith(marker))
        drop, seen = [], False
        for el in list(body):
            if seen and el.tag != qn("w:sectPr"):
                drop.append(el)
            if el is anchor:
                seen = True
        for el in drop:
            body.remove(el)

    def set_cover_line(self, startswith: str, text: str):
        p = next(p for p in self.doc.paragraphs if p.text.strip().startswith(startswith))
        runs = p.runs
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        return p

    def set_header(self, text: str):
        for sec in self.doc.sections:
            for p in sec.header.paragraphs:
                if p.text.strip():
                    p.runs[0].text = text
                    for r in p.runs[1:]:
                        r._r.getparent().remove(r._r)

    def update_fields_on_open(self):
        settings = self.doc.settings.element
        el = settings.find(qn("w:updateFields"))
        if el is None:
            el = OxmlElement("w:updateFields")
            settings.append(el)
        el.set(qn("w:val"), "true")

    # ---------- khối nội dung ----------
    def _runs(self, p, text: str, bold: bool = False):
        for part in _INLINE.split(text):
            if not part:
                continue
            if part.startswith("`"):
                r = p.add_run(part[1:-1])
                r.font.name = "Consolas"
                r.font.size = Pt(11)
            else:
                p.add_run(part).bold = bold or None
        return p

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def h1(self, text: str, chapter: bool = False):
        if chapter:
            self.chapter += 1
            self.n_table = self.n_fig = 0
        return self.doc.add_paragraph(text, style="Heading 1")

    def h2(self, text: str):
        return self.doc.add_paragraph(text, style="Heading 2")

    def h3(self, text: str):
        return self.doc.add_paragraph(text, style="Heading 3")

    def p(self, text: str, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False):
        par = self.doc.add_paragraph(style="Normal")
        par.alignment = align
        self._runs(par, text, bold)
        return par

    def bullets(self, items):
        for it in items:
            self._runs(self.doc.add_paragraph(style="List Bullet"), it)

    def code(self, text: str):
        par = self.doc.add_paragraph(style="CodeBlock")
        lines = text.strip("\n").split("\n")
        for i, line in enumerate(lines):
            r = par.add_run(line)
            if i < len(lines) - 1:
                r.add_break()
        return par

    def caption_table(self, title: str, prefix: str | None = None):
        self.n_table += 1
        num = prefix or f"{self.chapter}.{self.n_table}"
        return self.doc.add_paragraph(f"Bảng {num} {title}", style="CapBang")

    def table(self, title: str, headers, rows, widths_cm, prefix: str | None = None, center_cols=()):
        self.caption_table(title, prefix)
        total = sum(widths_cm)
        widths = [round(w / total * TEXT_WIDTH_DXA) for w in widths_cm]
        widths[-1] = TEXT_WIDTH_DXA - sum(widths[:-1])
        t = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.style = self.doc.styles["Table Grid"]
        t.alignment = 1  # center
        tblPr = t._tbl.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))
        for ri, row in enumerate([headers] + [list(r) for r in rows]):
            tr = t.rows[ri]
            if ri == 0:
                trPr = tr._tr.get_or_add_trPr()
                trPr.append(OxmlElement("w:tblHeader"))
            for ci, val in enumerate(row):
                cell = tr.cells[ci]
                tcPr = cell._tc.get_or_add_tcPr()
                tcW = OxmlElement("w:tcW")
                tcW.set(qn("w:w"), str(widths[ci]))
                tcW.set(qn("w:type"), "dxa")
                tcPr.append(tcW)
                if ri == 0:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), "D9E2F3")
                    tcPr.append(shd)
                par = cell.paragraphs[0]
                par.style = self.doc.styles["CellText"]
                if ri == 0 or ci in center_cols:
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._runs(par, str(val), ri == 0)
        self.doc.add_paragraph(style="Normal").paragraph_format.space_after = Pt(0)
        return t

    def figure(self, path: str, title: str, width_cm: float = 15.5):
        par = self.doc.add_paragraph(style="Normal")
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(path, width=Cm(width_cm))
        self.n_fig += 1
        return self.doc.add_paragraph(f"Hình {self.chapter}.{self.n_fig} {title}", style="CapHinh")

    def field(self, instr: str, placeholder: str, style: str = "Normal"):
        par = self.doc.add_paragraph(style=style)

        def fc(kind):
            r = OxmlElement("w:r")
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kind)
            r.append(f)
            return r

        par._p.append(fc("begin"))
        r = OxmlElement("w:r")
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = f" {instr} "
        r.append(it)
        par._p.append(r)
        par._p.append(fc("separate"))
        par.add_run(placeholder)
        par._p.append(fc("end"))
        return par

    def dotted_lines(self, n: int):
        for _ in range(n):
            self.p("." * 110, align=WD_ALIGN_PARAGRAPH.LEFT)

    def save(self, path: str):
        self.doc.save(path)
